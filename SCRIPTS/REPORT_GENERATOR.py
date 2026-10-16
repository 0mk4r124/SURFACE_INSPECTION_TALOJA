

#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
REPORT_GENERATOR.py   (v2)
SURFACE_INSPECTION_TALOJA -- white-spot validation against the site record.

WORKFLOW
  1. Site engineer fills THEIR_XLSX_PATH (the entry sheet). If that file does
     not exist yet, running this script creates a blank one there and stops.
  2. Set JSON_PATH to our coil JSON (history-table download).
  3. Run:   python3 REPORT_GENERATOR.py
  4. Open the result workbook (values only): SUMMARY, THEIR_SPOTS,
     MATCH_GROUPS, OUR_DETECTIONS. The summary is also printed.

MATCHING (v2)
  * Every detection in our JSON is kept - repeats are NOT merged.
  * Our metre -> their metre (their start = our end, coil unwinds):
        SCRAP_AT_OUR_END = True :
            their_m = (our_end - start_scrap - our_m) / K + OFFSET_M
        SCRAP_AT_OUR_END = False (old v1 rule):
            their_m = their_total - (our_m - start_scrap) / K + OFFSET_M
    our_end = our scanned length. Detections landing before their 0 m or past
    their total length are listed but not compared.
  * OS / DS is the same side on both systems. Our JSON gives pos_mm, measured
    from the centre of our 400 mm view: negative = OS, positive = DS. They
    measure from the edge of that side, so:
        side    = OS if pos_mm < 0, DS if pos_mm > 0
        edge_mm = coil_width / 2 - |pos_mm|
    Their spot is inside our view when |coil_width / 2 - edge_mm| <= 200 mm.
  * One of their spots is FOUND when at least one of our detections has
        the same side, |edge difference| <= TOL_LAT_MM, and
        |metre difference| <= max(TOL_LEN_M, STRETCH x their_m).
    All our detections that pass are listed against it (MATCH_GROUPS).
  * Two passes: pass 1 uses the formula with a window that grows with the
    stretch allowance; the spots it finds fit the real datum and stretch
    (their_m = a x our_m + b, outliers dropped); pass 2 re-matches everything
    on that fit within +- TOL_LEN_M.
  * Our model score for each detection (KEY_SCORE) is shown on every sheet, and
    SUMMARY compares the mean score of confirmed vs unconfirmed detections.
    MIN_SCORE (optional) leaves low-score detections out of the comparison.
  * Each of our detections is CONFIRMED (passes for some spot of theirs),
    NO COUNTERPART, NO OS/DS DATA, or IN SCRAP / OUT OF RANGE.

STANDALONE OFFLINE TOOL. Not a service, not started by MANAGE.py, imports
nothing from SCRIPTS/DEPENDANT. Takes no command-line arguments.
Requires: openpyxl
"""

import os
import re
import sys
import json
import bisect
import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.comments import Comment

# ============================================================================
# >>> USER SETTINGS <<<
# ============================================================================
JSON_PATH = "/home/hindalco_surface_inspection/Downloads/ROLL_20260921173744_defect_data.json"              # our history-table JSON
THEIR_XLSX_PATH = "/home/hindalco_surface_inspection/Downloads/XL_REPORT/WS_SITE_918.xlsx"          # site engineer's filled sheet
OUT_PATH = "/home/hindalco_surface_inspection/Downloads/XL_REPORT/"       # None = folder of the JSON | a folder | a full .xlsx path

# Metre alignment
SCRAP_AT_OUR_END = True   # their start scrap is cut from OUR end (their start)
K = 1.0                   # our metres per their metre (sheet stretch), e.g. 1.02
OFFSET_M = 0.0            # extra shift on their axis, after the formula

# Matching tolerances
TOL_LEN_M = 2.0           # pass 1: window near their 0 m, growing to STRETCH x metre;
STRETCH = 0.10            #         pass 2 (after the fit): fixed +- TOL_LEN_M
REFINE = True             # fit the real datum + stretch from pass-1 finds, re-match
MIN_FIT_SPOTS = 4         # found spots needed before the fit is trusted
TOL_LAT_MM = 40.0         # allowed distance-from-edge disagreement
TOL_W_MM = 5.0            # allowed white-spot width disagreement (reported only)
FRAME_WIDTH_MM = 400.0    # our camera field of view

# Our white-spot width = (box_px - 2 x BOX_PAD_PX) / PX_PER_MM.  None = no width.
PX_PER_MM = None
BOX_PAD_PX = None
BOX_FORMAT = "xyxy"       # "xyxy" = [x1,y1,x2,y2]   "xywh" = [x,y,w,h]
BOX_W_AXIS = "y"          # box axis that is the spot WIDTH

# Lateral position from the box, ONLY if the JSON has no x_mm-style field:
LAT_AXIS = None           # "x" or "y": image axis that runs across the coil
LAT_IMG_PX = None         # image size in px along LAT_AXIS (its centre = our centre)

# Per-defect entries nested inside each frame record (if the export has them)
LABEL_LIST_KEYS = ("labels", "defects", "dets", "defect_list", "boxes", "d")
LABEL_TUPLE = None        # if entries are lists, their field order,
                          # e.g. ("c", "x1", "y1", "x2", "y2", "score")

MIN_SCORE = None          # e.g. 0.5: our detections scoring below this are not compared
                          # (listed as BELOW MIN SCORE). None = compare all.

ENTRY_ROWS = 500          # rows in a newly created entry sheet

# ============================================================================
# JSON FIELD MAPPING - first key present in a record wins
# ============================================================================
KEY_LENGTH   = ("l", "length_mm", "len_mm", "position_mm", "meter_mm", "dist_mm",
                "length", "meter", "pos")
KEY_CLASS    = ("c", "cls", "class", "defect", "defect_type", "type", "label", "name")
KEY_XCENTER  = ("pos_mm", "x_mm", "x_center_mm", "xc_mm", "center_offset_mm",
                "offset_mm", "lat_mm", "dist_from_center_mm")
                # signed mm from the centre of our 400 mm view: - = OS, + = DS
KEY_SIDE     = ("side", "os_ds", "osds", "dsos")
KEY_BOX      = ("box", "bbox", "b", "rect", "xyxy")
KEY_BOX_W_PX = ("box_w_px", "bw_px", "bbox_w_px")
KEY_FRAME    = ("r", "frame", "frame_no", "frame_index", "frame_id", "idx")
KEY_IMAGE    = ("img", "image", "image_path", "path", "file", "filename")
KEY_SCORE    = ("score", "conf", "confidence", "accuracy", "acc", "prob", "probability")
                # our model's score per white spot (a value, or a list like pos_mm)
KEY_TYPEMAP  = ("t", "types", "classes", "defect_counts", "counts")
KEY_RESULT_LIST = ("result", "results", "defects", "defect_details",
                   "details", "data", "rows", "records")
WS_CLASS_CODES = {"W", "V", "WS", "WHITE_SPOT", "WHITESPOT", "WHITE SPOT"}
JSON_LENGTH_UNIT = "mm"
POSITIVE_SIDE = "DS"      # pos_mm > 0 is DS, pos_mm < 0 is OS
CENTRE_BAND_MM = 5.0     # |pos_mm| below this sits on the centre line: matches OS or DS

# ============================================================================
# STYLE
# ============================================================================
FONT = "Arial"
F_TITLE = Font(name=FONT, size=13, bold=True, color="FFFFFF")
F_HDR = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_LAB = Font(name=FONT, size=10, bold=True)
F_BODY = Font(name=FONT, size=10)
F_IN = Font(name=FONT, size=10, color="0000FF")
F_NOTE = Font(name=FONT, size=9, italic=True, color="808080")
F_BIG = Font(name=FONT, size=12, bold=True)
FILL_TITLE = PatternFill("solid", fgColor="1F3864")
FILL_HDR = PatternFill("solid", fgColor="2F5597")
FILL_IN = PatternFill("solid", fgColor="FFFF00")
FILL_KEY = PatternFill("solid", fgColor="E2EFDA")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_BAD = PatternFill("solid", fgColor="FFC7CE")
FILL_MID = PatternFill("solid", fgColor="FFEB9C")
FILL_GREY = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CTR = Alignment(horizontal="center", vertical="center", wrap_text=True)

STATUS_FILL = {"FOUND": FILL_OK, "CONFIRMED": FILL_OK, "FOUND (NO LATERAL)": FILL_MID,
               "SIZE MISMATCH": FILL_MID, "LENGTH ONLY": FILL_MID,
               "MISSED BY US": FILL_BAD, "NO COUNTERPART": FILL_BAD,
               "OUTSIDE OUR VIEW": FILL_GREY, "IN SCRAP / OUT OF RANGE": FILL_GREY,
               "NO OS/DS DATA": FILL_GREY, "BELOW MIN SCORE": FILL_GREY}


# ============================================================================
# 1. SITE ENTRY SHEET (template + reader)
# ============================================================================
ENTRY_LABELS = [("Coil ID", "coil_id", None),
                ("Total coil length (m)", "total_len_m", "0.000"),
                ("Start scrap length (m)", "scrap_m", "0.000"),
                ("Coil width (mm)", "coil_w_mm", "0")]
ENTRY_COLS = [("SR", 6), ("SIDE\n(OS / DS)", 11), ("DEFECT WIDTH\n(mm)", 14),
              ("DISTANCE FROM\nSIDE EDGE (mm)", 17), ("LENGTH FOUND\nAT (m)", 15),
              ("REMARKS", 30)]


def make_entry_template(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "SITE_ENTRY"
    ws.sheet_view.showGridLines = False
    ws.merge_cells("A1:F1")
    ws["A1"] = "WHITE SPOT RECORD - SITE ENGINEER ENTRY"
    ws["A1"].font = F_TITLE
    ws["A1"].fill = FILL_TITLE
    ws.row_dimensions[1].height = 24
    for i, (lab, _, fmt) in enumerate(ENTRY_LABELS):
        r = 3 + i
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=2)
        ws.cell(r, 1, lab).font = F_LAB
        c = ws.cell(r, 3)
        c.fill = FILL_IN
        c.font = F_IN
        c.border = BOX
        if fmt:
            c.number_format = fmt
    ws.merge_cells("D3:F6")
    ws["D3"] = ("One row per white spot. Rows can be in any order. SIDE is the edge "
                "the distance is measured from. Leave blank rows anywhere - they are "
                "ignored. Only LENGTH FOUND AT is required on a row.")
    ws["D3"].font = F_NOTE
    ws["D3"].alignment = Alignment(wrap_text=True, vertical="top")
    hr = 8
    for i, (h, w) in enumerate(ENTRY_COLS, start=1):
        c = ws.cell(hr, i, h)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = CTR
        c.border = BOX
        ws.column_dimensions[chr(64 + i)].width = w
    ws.row_dimensions[hr].height = 32
    ws["E8"].comment = Comment("Metre on YOUR axis where the spot was found.", "REPORT_GENERATOR")
    ws["D8"].comment = Comment("Distance of the spot from the coil edge on the side "
                               "given in column B, in mm.", "REPORT_GENERATOR")
    for r in range(hr + 1, hr + 1 + ENTRY_ROWS):
        ws.cell(r, 1, r - hr).font = F_NOTE
        ws.cell(r, 1).alignment = CTR
        for c in range(2, 7):
            cell = ws.cell(r, c)
            cell.fill = FILL_IN
            cell.font = F_IN
            cell.border = BOX
        ws.cell(r, 3).number_format = "0.0"
        ws.cell(r, 4).number_format = "0.0"
        ws.cell(r, 5).number_format = "0.000"
    dv = DataValidation(type="list", formula1='"OS,DS"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add("B%d:B%d" % (hr + 1, hr + ENTRY_ROWS))
    ws.freeze_panes = "A9"

    ex = wb.create_sheet("EXAMPLE")
    ex.sheet_view.showGridLines = False
    ex["A1"] = "How a filled sheet looks (this tab is not read)"
    ex["A1"].font = F_LAB
    rows = [("Coil ID", "26LPHU020944"), ("Total coil length (m)", 2958),
            ("Start scrap length (m)", 12), ("Coil width (mm)", 1250)]
    for i, (a, b) in enumerate(rows):
        ex.cell(3 + i, 1, a).font = F_BODY
        ex.cell(3 + i, 3, b).font = F_BODY
    for i, (h, w) in enumerate(ENTRY_COLS, start=1):
        c = ex.cell(8, i, h)
        c.font = F_HDR
        c.fill = FILL_HDR
        c.alignment = CTR
        ex.column_dimensions[chr(64 + i)].width = w
    ex.row_dimensions[8].height = 32
    for i, row in enumerate([(1, "DS", 12, 310, 41.5, ""), (2, "OS", 8, 455, 47.2, ""),
                             (3, "DS", 15, 380, 55.0, "faint")], start=9):
        for c, v in enumerate(row, start=1):
            ex.cell(i, c, v).font = F_BODY
    os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
    wb.save(path)


def _f(v):
    if v is None or isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.match(r"^\s*[-+]?\d*\.?\d+", str(v).replace(",", ""))
    return float(m.group(0)) if m else None


def _side(v):
    s = str(v).strip().upper() if v is not None else ""
    return "OS" if s.startswith("O") else ("DS" if s.startswith("D") else None)


def read_entry(path):
    """Read the site sheet by labels / header text, so moved rows still work."""
    wb = load_workbook(path, data_only=True)
    ws = wb["SITE_ENTRY"] if "SITE_ENTRY" in wb.sheetnames else wb.worksheets[0]
    info, warns = {}, []
    hdr_row, cols = None, {}
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if isinstance(a, str):
            low = a.strip().lower()
            for lab, key, _ in ENTRY_LABELS:
                if low.startswith(lab.lower().split(" (")[0]):
                    for c in range(2, 6):
                        v = ws.cell(r, c).value
                        if v not in (None, ""):
                            info[key] = v
                            break
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if isinstance(v, str) and v.strip().upper().startswith("LENGTH FOUND"):
                hdr_row = r
        if hdr_row:
            break
    if not hdr_row:
        sys.exit("Could not find the 'LENGTH FOUND AT (m)' header in %s" % path)
    for c in range(1, ws.max_column + 1):
        h = str(ws.cell(hdr_row, c).value or "").upper()
        if h.startswith("SR"):
            cols["sr"] = c
        elif "SIDE" in h and "DISTANCE" not in h:
            cols["side"] = c
        elif "WIDTH" in h:
            cols["w"] = c
        elif "DISTANCE" in h:
            cols["edge"] = c
        elif "LENGTH" in h:
            cols["len"] = c
        elif "REMARK" in h:
            cols["rem"] = c
    spots = []
    for r in range(hdr_row + 1, ws.max_row + 1):
        L = _f(ws.cell(r, cols["len"]).value)
        if L is None:
            continue
        g = lambda k: ws.cell(r, cols[k]).value if k in cols else None
        spots.append(dict(sr=g("sr") or (r - hdr_row), m=L, side=_side(g("side")),
                          edge=_f(g("edge")), w=_f(g("w")), rem=g("rem")))
    for key, lab in (("total_len_m", "Total coil length"), ("coil_w_mm", "Coil width")):
        if _f(info.get(key)) is None:
            warns.append("%s is empty on the site sheet." % lab)
    return info, spots, warns



# ============================================================================
# 2. OUR JSON
# ============================================================================
def _pick(rec, keys):
    for k in keys:
        if k in rec and rec[k] not in (None, ""):
            return rec[k]
    return None


def _is_ws(code):
    return code is not None and str(code).strip().upper().replace("-", "_") in WS_CLASS_CODES


def _box(rec):
    """(x1, y1, x2, y2) in px from any supported box form, else None."""
    b = _pick(rec, KEY_BOX)
    if b is None and all(k in rec for k in ("x1", "y1", "x2", "y2")):
        b = [rec["x1"], rec["y1"], rec["x2"], rec["y2"]]
    if b is None:
        return None
    if isinstance(b, str):
        b = [float(x) for x in re.findall(r"[-+]?\d*\.?\d+", b)]
    if isinstance(b, dict):
        g = lambda *ks: next((_f(b[k]) for k in ks if k in b), None)
        x1, y1 = g("x1", "xmin", "x"), g("y1", "ymin", "y")
        x2, y2 = g("x2", "xmax"), g("y2", "ymax")
        w, h = g("w", "width"), g("h", "height")
        if x2 is None and None not in (x1, w):
            x2 = x1 + w
        if y2 is None and None not in (y1, h):
            y2 = y1 + h
        v = [x1, y1, x2, y2]
    elif isinstance(b, (list, tuple)) and len(b) >= 4:
        v = [_f(x) for x in b[:4]]
        if None not in v and BOX_FORMAT == "xywh":
            v = [v[0], v[1], v[0] + v[2], v[1] + v[3]]
    else:
        return None
    return None if None in v else v


def _scalar(v):
    """A single number (also "0.87"), or None for a list / list-like string."""
    if isinstance(v, (list, tuple)) or (isinstance(v, str) and v.strip().startswith("[")):
        return None
    return _f(v)


def _set_lat(s, x):
    """pos_mm (signed, from our centre) -> side + unsigned distance from centre."""
    if x is None:
        return
    if abs(x) < CENTRE_BAND_MM:
        s["side"] = "C"                                   # on the centre line
    else:
        s["side"] = POSITIVE_SIDE if x > 0 else ("OS" if POSITIVE_SIDE == "DS" else "DS")
    s["dist"] = abs(x)
    s["pos"] = x


def _defect(item, base):
    """One per-defect dict -> our spot dict, or None if not a white spot."""
    if not _is_ws(_pick(item, KEY_CLASS)):
        return None
    s = dict(base)
    L = _f(_pick(item, KEY_LENGTH))
    if L is not None:
        s["our_m"] = L / (1000.0 if JSON_LENGTH_UNIT == "mm" else 1.0)
    s["img"] = _pick(item, KEY_IMAGE) or base["img"]
    box = _box(item)
    x = _f(_pick(item, KEY_XCENTER))
    if x is None and box and LAT_AXIS and LAT_IMG_PX and PX_PER_MM:
        c = (box[0] + box[2]) / 2.0 if LAT_AXIS == "x" else (box[1] + box[3]) / 2.0
        x = (c - LAT_IMG_PX / 2.0) / PX_PER_MM
    _set_lat(s, x)
    sc = _f(_pick(item, KEY_SCORE))
    if sc is not None:
        s["score"] = sc
    side = _side(_pick(item, KEY_SIDE))           # an explicit side field wins
    if side:
        s["side"] = side
    bw = _f(_pick(item, KEY_BOX_W_PX))
    if bw is None and box:
        bw = abs(box[3] - box[1]) if BOX_W_AXIS == "y" else abs(box[2] - box[0])
    s["box"] = bw
    return s


def _records(payload):
    if isinstance(payload, list):
        return payload
    if not isinstance(payload, dict):
        return []
    for k in KEY_RESULT_LIST:
        v = payload.get(k)
        if isinstance(v, str):
            try:
                v = json.loads(v)
            except (ValueError, TypeError):
                continue
        if isinstance(v, list):
            return v
        if isinstance(v, dict):
            return _records(v)
    best = []
    for v in payload.values():
        if isinstance(v, list) and v and isinstance(v[0], dict) and len(v) > len(best):
            best = v
    return best


def read_json(path):
    with open(path, "r") as fh:
        payload = json.load(fh)
    meta = payload if isinstance(payload, dict) else {}
    unit = 1000.0 if JSON_LENGTH_UNIT == "mm" else 1.0
    spots, warns, frame_level, bad_list = [], [], 0, 0
    recs = _records(payload)
    first = next((r for r in recs if isinstance(r, dict)), None)
    for rec in recs:
        if not isinstance(rec, dict):
            continue
        L = _f(_pick(rec, KEY_LENGTH))
        if L is None:
            continue
        base = dict(our_m=L / unit, img=_pick(rec, KEY_IMAGE),
                    side=None, dist=None, pos=None, box=None,
                    score=_scalar(_pick(rec, KEY_SCORE)))
        nested = None
        for k in LABEL_LIST_KEYS:
            v = rec.get(k)
            if isinstance(v, list) and v and isinstance(v[0], (dict, list, tuple)):
                nested = v
                break
        if nested is not None:                       # per-defect list inside the frame
            for item in nested:
                if isinstance(item, (list, tuple)):
                    if not LABEL_TUPLE:
                        continue
                    item = dict(zip(LABEL_TUPLE, item))
                d = _defect(item, base) if isinstance(item, dict) else None
                if d:
                    spots.append(d)
            continue
        tmap = _pick(rec, KEY_TYPEMAP)
        if isinstance(tmap, dict):                   # frame record with class counts
            lat = _pick(rec, KEY_XCENTER)
            if isinstance(lat, str) and lat.strip().startswith("["):
                try:
                    lat = json.loads(lat)
                except ValueError:
                    pass
            ws_n = 0
            for code, cnt in tmap.items():
                if _is_ws(code):
                    try:
                        ws_n += max(int(cnt), 1)
                    except (TypeError, ValueError):
                        ws_n += 1
            if not ws_n:
                continue
            if isinstance(lat, (list, tuple)):
                vals = [_f(v) for v in lat]
                if len(vals) != ws_n:                # cannot tell which value is which
                    bad_list += 1
                    vals = [None] * ws_n
            else:
                vals = [_f(lat)] * ws_n              # one position for the frame
            sc = _pick(rec, KEY_SCORE)
            if isinstance(sc, str) and sc.strip().startswith("["):
                try:
                    sc = json.loads(sc)
                except ValueError:
                    pass
            if isinstance(sc, (list, tuple)):
                svals = [_f(v) for v in sc]
                if len(svals) != ws_n:
                    bad_list += 1
                    svals = [None] * ws_n
            else:
                svals = [_f(sc)] * ws_n
            for x, sv in zip(vals, svals):
                sp = dict(base)
                sp["score"] = sv
                _set_lat(sp, x)
                if x is None:
                    frame_level += 1
                spots.append(sp)
            continue
        d = _defect(rec, base)                       # record is itself one defect
        if d:
            spots.append(d)
    if not recs:
        warns.append("No defect list found in the JSON - check KEY_RESULT_LIST.")
    elif not spots:
        warns.append("No white spots in the JSON - check WS_CLASS_CODES.")
    with_lat = sum(1 for s in spots if s["side"] and s["dist"] is not None)
    if bad_list:
        warns.append("%d frame record(s) had a pos_mm or score list whose length did not "
                     "match the white-spot count - those values were skipped." % bad_list)
    if spots and not any(s.get("score") is not None for s in spots):
        warns.append("No score found for any white spot - add the JSON key to KEY_SCORE. "
                     "First record: %s" % (json.dumps(first, default=str)[:300] if first else "-"))
    wide = sum(1 for s in spots if s["dist"] is not None and s["dist"] > FRAME_WIDTH_MM / 2 + 5)
    if wide:
        warns.append("%d detection(s) have |pos_mm| beyond %g mm (outside our view) - check "
                     "that pos_mm is in mm." % (wide, FRAME_WIDTH_MM / 2))
    if spots and with_lat == 0:
        sample = json.dumps(first, default=str)[:400] if first else "-"
        warns.append("NO pos_mm FOUND FOR ANY WHITE SPOT (%d detections). OS/DS "
                     "validation is impossible: their spots can only reach LENGTH ONLY. "
                     "Add the real key name to KEY_XCENTER. First record: %s"
                     % (frame_level, sample))
    elif spots and with_lat < len(spots):
        warns.append("%d of %d our detections have no OS/DS and are not compared."
                     % (len(spots) - with_lat, len(spots)))
    meta["_our_end_m"] = ((_f(meta.get("coil_length_mm")) or 0) / 1000.0 or
                          max((s["our_m"] for s in spots), default=0.0))
    meta["_first_record"] = first
    return meta, spots, warns


# ============================================================================
# 3. MATCHING
# ============================================================================
def to_their(our_m, our_end, total, scrap):
    if SCRAP_AT_OUR_END:
        return (our_end - scrap - our_m) / K + OFFSET_M
    return total - (our_m - scrap) / K + OFFSET_M


def fit_line(xs, ys):
    n = len(xs)
    if n < 2:
        return None, None
    mx, my = sum(xs) / n, sum(ys) / n
    sxx = sum((x - mx) ** 2 for x in xs)
    if sxx == 0:
        return None, None
    a = sum((x - mx) * (y - my) for x, y in zip(xs, ys)) / sxx
    return a, my - a * mx


# ============================================================================
# 4. RESULT WORKBOOK (values only)
# ============================================================================
def _col(n):
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def write_table(ws, r0, headers, rows, fmts=None, status_col=None, widths=None, light=False):
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(r0, c, h)
        cell.font = F_HDR
        cell.fill = FILL_HDR
        cell.alignment = CTR
        cell.border = BOX
    ws.row_dimensions[r0].height = 30
    for i, row in enumerate(rows, start=r0 + 1):
        for c, v in enumerate(row, start=1):
            cell = ws.cell(i, c, v)
            if not light:
                cell.font = F_BODY
                cell.border = BOX
            if fmts and fmts[c - 1]:
                cell.number_format = fmts[c - 1]
            if status_col == c and v in STATUS_FILL:
                cell.fill = STATUS_FILL[v]
    if widths:
        for c, w in enumerate(widths, start=1):
            ws.column_dimensions[_col(c)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1)
    if rows:
        ws.auto_filter.ref = "A%d:%s%d" % (r0, _col(len(headers)), r0 + len(rows))


def title(ws, text, span="A1:H1"):
    ws.merge_cells(span)
    a = span.split(":")[0]
    ws[a] = text
    ws[a].font = F_TITLE
    ws[a].fill = FILL_TITLE
    ws.row_dimensions[1].height = 22
    ws.sheet_view.showGridLines = False


def rnd(v, n):
    return round(v, n) if isinstance(v, (int, float)) else v


def main():
    # ---------------------------------------------------------------- inputs
    if not os.path.isfile(THEIR_XLSX_PATH):
        make_entry_template(THEIR_XLSX_PATH)
        print("Blank site entry sheet created:\n  %s\nFill it, then run again."
              % THEIR_XLSX_PATH)
        return
    if not os.path.isfile(JSON_PATH):
        sys.exit("JSON_PATH not found: %s" % JSON_PATH)

    meta, ours, warns = read_json(JSON_PATH)
    info, theirs, w2 = read_entry(THEIR_XLSX_PATH)
    warns += w2
    our_end = meta["_our_end_m"]
    total = _f(info.get("total_len_m"))
    scrap = _f(info.get("scrap_m")) or 0.0
    coil_w = _f(info.get("coil_w_mm"))
    if total is None:
        total = our_end - scrap
        warns.append("Total coil length missing - used our length minus scrap (%.3f m)." % total)
    jid, sid = str(meta.get("coil_id") or ""), str(info.get("coil_id") or "")
    if jid and sid and jid.strip().upper() != sid.strip().upper():
        warns.append("Coil ID differs: JSON '%s' vs site sheet '%s'." % (jid, sid))
    if not coil_w:
        warns.append("Coil width missing - edge distances cannot be compared.")

    # ---------------------------------------------------------------- our side
    for i, o in enumerate(ours):
        o["id"] = i + 1
        o["edge"] = (coil_w / 2.0 - o["dist"]) if (coil_w and o["dist"] is not None) else None
        o["has_lat"] = bool(o["side"]) and o["edge"] is not None
        o["low"] = (MIN_SCORE is not None and o.get("score") is not None
                    and o["score"] < MIN_SCORE)
        o["w"] = (max(0.0, o["box"] - 2 * BOX_PAD_PX) / PX_PER_MM
                  if (o["box"] is not None and PX_PER_MM and BOX_PAD_PX is not None) else None)

    def set_axis(fn):
        for o in ours:
            o["t"] = fn(o["our_m"])
            o["in_range"] = -TOL_LEN_M <= o["t"] <= total + TOL_LEN_M
            o["srs"] = []
        ours.sort(key=lambda o: o["t"])
        return [o["t"] for o in ours]

    # ---------------------------------------------------------------- their side
    judged, outside = [], []
    for s in theirs:
        s["dist"] = (coil_w / 2.0 - s["edge"]) if (coil_w and s["edge"] is not None) else None
        if s["dist"] is not None and abs(s["dist"]) > FRAME_WIDTH_MM / 2.0:
            s["fov"] = "NO"
            outside.append(s)
        else:
            s["fov"] = "YES" if s["dist"] is not None else "UNKNOWN"
            judged.append(s)

    # ---------------------------------------------------------------- match
    def run(tkeys, tol_fn):
        their_rows, group_rows, pairs = [], [], []
        found = length_only = 0
        for s in sorted(judged, key=lambda s: s["m"]):
            tol = tol_fn(s["m"])
            lo = bisect.bisect_left(tkeys, s["m"] - tol)
            hi = bisect.bisect_right(tkeys, s["m"] + tol)
            window = [o for o in ours[lo:hi] if o["in_range"] and not o["low"]]
            valid = []
            for o in window:
                if not o["has_lat"]:
                    continue
                if s["side"] and o["side"] not in (s["side"], "C"):
                    continue
                if s["edge"] is not None and abs(o["edge"] - s["edge"]) > TOL_LAT_MM:
                    continue
                valid.append(o)
            best = None
            if valid:
                def cost(o):
                    de = (o["edge"] - s["edge"]) / TOL_LAT_MM if s["edge"] is not None else 0.0
                    return ((o["t"] - s["m"]) / tol) ** 2 + de ** 2
                best = min(valid, key=cost)
                status = "FOUND" if s["edge"] is not None else "FOUND (NO LATERAL)"
                found += 1
                pairs.append((best["our_m"], s["m"]))
                for o in valid:
                    o["srs"].append(s["sr"])
                    group_rows.append([s["sr"], rnd(s["m"], 3), s["side"], rnd(s["edge"], 1),
                                       rnd(s["w"], 1), o["id"], rnd(o["our_m"], 3),
                                       rnd(o["t"], 3), rnd(o["t"] - s["m"], 3), o["side"],
                                       rnd(o["edge"], 1),
                                       rnd(o["edge"] - s["edge"], 1) if s["edge"] is not None else None,
                                       rnd(o["w"], 1), rnd(o.get("score"), 4),
                                       "BEST" if o is best else None,
                                       os.path.basename(str(o["img"])) if o["img"] else None])
            elif any(not o["has_lat"] for o in window):
                status = "LENGTH ONLY"            # nearby, but no OS/DS to check
                length_only += 1
            else:
                status = "MISSED BY US"           # nothing near, or wrong side / lateral
            dw = (best["w"] - s["w"]) if (best and best["w"] is not None and s["w"] is not None) else None
            their_rows.append([
                s["sr"], rnd(s["m"], 3), s["side"], rnd(s["edge"], 1), rnd(s["w"], 1), s["fov"],
                rnd(tol, 2), len(window), sum(1 for o in window if o["has_lat"]), len(valid),
                status, rnd(best["our_m"], 3) if best else None, rnd(best["t"], 3) if best else None,
                rnd(best["t"] - s["m"], 3) if best else None,
                best["side"] if best else None, rnd(best["edge"], 1) if best else None,
                rnd(best["edge"] - s["edge"], 1) if (best and s["edge"] is not None) else None,
                rnd(best["w"], 1) if best else None, rnd(dw, 1),
                ("OK" if abs(dw) <= TOL_W_MM else "OUT") if dw is not None else None,
                rnd(best.get("score"), 4) if best else None,
                os.path.basename(str(best["img"])) if (best and best["img"]) else None])
        return their_rows, group_rows, pairs, found, length_only

    # pass 1: formula alignment, window grows with stretch
    tkeys = set_axis(lambda m: to_their(m, our_end, total, scrap))
    their_rows, group_rows, pairs, found, length_only = run(
        tkeys, lambda m: max(TOL_LEN_M, STRETCH * abs(m)))
    align = "formula only (%d spots found - need %d to refine)" % (found, MIN_FIT_SPOTS)
    fit_used = None
    # pass 2: fit their_m = a x our_m + b on the found spots, drop outliers, re-match tight
    if REFINE and len(pairs) >= MIN_FIT_SPOTS:
        pts = list(pairs)
        for _ in range(3):
            a, b = fit_line([p[0] for p in pts], [p[1] for p in pts])
            if a is None:
                break
            res = [abs(y - (a * x + b)) for x, y in pts]
            med = sorted(res)[len(res) // 2]
            keep = [p for p, r in zip(pts, res) if r <= max(TOL_LEN_M, 3 * med)]
            if len(keep) == len(pts) or len(keep) < MIN_FIT_SPOTS:
                break
            pts = keep
        if a is not None and a < 0 and 0.8 <= -1.0 / a <= 1.25:
            fit_used = (a, b, len(pts), len(pairs))
            tkeys = set_axis(lambda m: a * m + b)
            their_rows, group_rows, pairs, found, length_only = run(tkeys, lambda m: TOL_LEN_M)
            align = ("fitted on %d of %d spots, re-matched within +-%g m"
                     % (len(pts), fit_used[3], TOL_LEN_M))
        else:
            align = "fit rejected (implausible K) - formula alignment kept"
    for s in outside:
        their_rows.append([s["sr"], rnd(s["m"], 3), s["side"], rnd(s["edge"], 1), rnd(s["w"], 1),
                           "NO"] + [None] * 4 + ["OUTSIDE OUR VIEW"] + [None] * 11)

    our_rows, n_cmp, n_conf, n_nolat, n_out, n_low = [], 0, 0, 0, 0, 0
    sc_conf, sc_none, sc_cmp = [], [], []
    for o in ours:
        if not o["in_range"]:
            st = "IN SCRAP / OUT OF RANGE"
            n_out += 1
        elif o["low"]:
            st = "BELOW MIN SCORE"
            n_low += 1
        elif not o["has_lat"]:
            st = "NO OS/DS DATA"
            n_nolat += 1
        elif o["srs"]:
            st = "CONFIRMED"
            n_cmp += 1
            n_conf += 1
        else:
            st = "NO COUNTERPART"
            n_cmp += 1
        if o.get("score") is not None and st in ("CONFIRMED", "NO COUNTERPART"):
            sc_cmp.append(o["score"])
            (sc_conf if st == "CONFIRMED" else sc_none).append(o["score"])
        our_rows.append([o["id"], rnd(o["our_m"], 3), rnd(o["t"], 3), rnd(o["pos"], 1),
                         o["side"], rnd(o["edge"], 1), rnd(o["w"], 1), rnd(o.get("score"), 4), st,
                         ", ".join(str(x) for x in o["srs"]) or None,
                         os.path.basename(str(o["img"])) if o["img"] else None])

    # ---------------------------------------------------------------- metrics
    nT = len(judged)
    all_sc = [o["score"] for o in ours if o.get("score") is not None]
    SC_FMT = "0.0%" if (all_sc and max(all_sc) <= 1.0) else "0.00"
    mean = lambda v: (sum(v) / len(v)) if v else None
    best_sc = [r[20] for r in their_rows if r[10] in ("FOUND", "FOUND (NO LATERAL)")
               and r[20] is not None]
    rec = found / float(nT) if nT else None
    prec = n_conf / float(n_cmp) if n_cmp else None
    f1 = (2 * prec * rec / (prec + rec)) if (prec and rec) else (0.0 if (prec is not None and rec is not None) else None)
    if fit_used:
        a_fit, b_fit = fit_used[0], fit_used[1]
    else:
        a_fit, b_fit = fit_line([p[0] for p in pairs], [p[1] for p in pairs])
    k_fit = (-1.0 / a_fit) if a_fit else None
    zero_fit = (-b_fit / a_fit) if a_fit else None

    # ---------------------------------------------------------------- output path
    uid = meta.get("cycle_uid") or os.path.splitext(os.path.basename(JSON_PATH))[0]
    fname = "WS_RESULT_%s.xlsx" % uid
    out = os.path.expanduser(OUT_PATH) if OUT_PATH else ""
    if not out:
        out = os.path.join(os.path.dirname(os.path.abspath(JSON_PATH)), fname)
    elif os.path.isdir(out) or not out.lower().endswith(".xlsx"):
        out = os.path.join(out, fname)
    os.makedirs(os.path.dirname(os.path.abspath(out)), exist_ok=True)

    wb = Workbook()
    sm = wb.active
    sm.title = "SUMMARY"
    title(sm, "WHITE SPOT VALIDATION - RESULT", "A1:C1")
    sm.column_dimensions["A"].width = 50
    sm.column_dimensions["B"].width = 18
    sm.column_dimensions["C"].width = 66
    pct = "0.0%"
    lines = [
        ("SPOT DETECTION ACCURACY", None, None, None),
        ("Spot detection accuracy (their spots we found)", rec, pct,
         "FOUND / their spots judged. Found = same side, edge within %g mm, metre in window."
         % TOL_LAT_MM),
        ("Detection confirmation (our detections they confirm)", prec, pct,
         "CONFIRMED / our detections compared (in range, with OS/DS)."),
        ("F1 score", f1, pct, None),
        ("THEIR SPOTS", None, None, None),
        ("Entries on the site sheet", len(theirs), "0", None),
        ("Outside our %d mm view (excluded)" % FRAME_WIDTH_MM, len(outside), "0",
         "Edge blind band - we cannot see it."),
        ("Judged", nT, "0", None),
        ("FOUND", found, "0", None),
        ("LENGTH ONLY (our detections nearby, but no OS/DS to check)", length_only, "0",
         "Not counted as found."),
        ("MISSED BY US", nT - found - length_only, "0", None),
        ("OUR DETECTIONS (none merged)", None, None, None),
        ("Detections in JSON", len(ours), "0", None),
        ("In their scrap / outside their length (not compared)", n_out, "0", None),
        ("No OS/DS in JSON (not compared)", n_nolat, "0", None),
        ("Compared", n_cmp, "0", None),
        ("CONFIRMED", n_conf, "0", "Passes the tests for at least one of their spots."),
        ("NO COUNTERPART", n_cmp - n_conf, "0", None),
        ("Below MIN_SCORE (not compared)", n_low, "0",
         "MIN_SCORE = %s" % MIN_SCORE),
        ("OUR DETECTION SCORE (from JSON)", None, None, None),
        ("Detections with a score", len(all_sc), "0", None),
        ("Mean score - CONFIRMED detections", mean(sc_conf), SC_FMT,
         "Detections the site record agrees with."),
        ("Mean score - NO COUNTERPART detections", mean(sc_none), SC_FMT,
         "Much lower than CONFIRMED = a score threshold would remove false detections."),
        ("Mean score - best match of each found spot", mean(best_sc), SC_FMT, None),
        ("Lowest score among those best matches", min(best_sc) if best_sc else None, SC_FMT,
         "A MIN_SCORE above this would start losing real spots."),
        ("ALIGNMENT", None, None, None),
        ("Rule", "scrap at our end" if SCRAP_AT_OUR_END else "scrap at our start", None,
         "K = %g, OFFSET_M = %g" % (K, OFFSET_M)),
        ("Alignment used", align, None, None),
        ("Their 0 m = our metre", rnd(our_end - scrap + OFFSET_M * K if SCRAP_AT_OUR_END
                                      else scrap + (total + OFFSET_M) * K, 3), "0.000",
         "Our scanned length %.3f m." % our_end),
        ("Fitted: their 0 m = our metre", zero_fit, "0.000",
         "Compare with the row above. A big gap means the scrap / datum rule is off."),
        ("Fitted: K (our m per their m, i.e. stretch)", k_fit, "0.0000", None),
        ("COIL", None, None, None),
        ("Coil ID (JSON / site)", "%s / %s" % (jid or "-", sid or "-"), None, None),
        ("Their total length / start scrap (m)", "%s / %s" % (total, scrap), None, None),
        ("Coil width (mm)", coil_w, "0", None),
        ("Tolerances", "%g m + %g%% / %g mm / %g mm" % (TOL_LEN_M, STRETCH * 100,
                                                     TOL_LAT_MM, TOL_W_MM), None,
         "metre window / edge distance / width"),
        ("Generated", datetime.datetime.now().strftime("%Y-%m-%d %H:%M"), None,
         "%s  +  %s" % (os.path.basename(JSON_PATH), os.path.basename(THEIR_XLSX_PATH))),
    ]
    r = 3
    for lab, val, fmt, note in lines:
        if val is None and fmt is None and note is None:
            r += 1
            sm.cell(r, 1, lab).font = Font(name=FONT, size=10, bold=True, color="1F3864")
            r += 1
            continue
        big = lab.startswith("Spot detection accuracy")
        sm.cell(r, 1, lab).font = F_BIG if big else F_BODY
        c = sm.cell(r, 2, val)
        c.font = F_BIG if big else F_BODY
        c.alignment = CTR
        c.border = BOX
        if fmt:
            c.number_format = fmt
        if big:
            c.fill = FILL_KEY
            sm.cell(r, 1).fill = FILL_KEY
        if note:
            sm.cell(r, 3, note).font = F_NOTE
        r += 1
    if warns:
        r += 1
        sm.cell(r, 1, "WARNINGS").font = Font(name=FONT, size=10, bold=True, color="C00000")
        for w in warns:
            r += 1
            sm.merge_cells(start_row=r, start_column=1, end_row=r, end_column=3)
            c = sm.cell(r, 1, "! " + w)
            c.font = Font(name=FONT, size=9, color="833C00")
            c.alignment = Alignment(wrap_text=True, vertical="top")
            sm.row_dimensions[r].height = 15 * (len(w) // 150 + 1)

    ws = wb.create_sheet("THEIR_SPOTS")
    title(ws, "THEIR SPOTS - validated against all our detections (their metre axis)", "A1:V1")
    write_table(ws, 3, ["SR", "THEIR_M", "SIDE", "EDGE_MM", "W_MM", "IN_OUR_VIEW",
                        "WINDOW_+-M", "OURS_IN_WINDOW", "WITH_OS/DS", "MATCHING",
                        "STATUS", "BEST_OUR_M", "BEST_ON_THEIR_AXIS", "D_LEN_M", "BEST_SIDE",
                        "BEST_EDGE_MM", "D_EDGE_MM", "BEST_W_MM", "D_W_MM", "W_CHECK",
                        "BEST_SCORE", "BEST_IMAGE"], their_rows,
                [None, "0.000", None, "0.0", "0.0", None, "0.00", None, None, None, None,
                 "0.000", "0.000", "0.000", None, "0.0", "0.0", "0.0", "0.0", None, SC_FMT,
                 None], 11,
                [6, 10, 6, 9, 7, 9, 10, 10, 9, 9, 17, 11, 12, 9, 8, 10, 9, 9, 8, 8, 10, 55])

    ws = wb.create_sheet("MATCH_GROUPS")
    title(ws, "MATCH GROUPS - every one of our detections that validates each of their spots",
          "A1:P1")
    write_table(ws, 3, ["THEIR_SR", "THEIR_M", "SIDE", "THEIR_EDGE", "THEIR_W", "OUR_ID",
                        "OUR_M", "ON_THEIR_AXIS", "D_LEN_M", "OUR_SIDE", "OUR_EDGE",
                        "D_EDGE_MM", "OUR_W", "OUR_SCORE", "BEST", "IMAGE"], group_rows,
                [None, "0.000", None, "0.0", "0.0", None, "0.000", "0.000", "0.000", None,
                 "0.0", "0.0", "0.0", SC_FMT, None, None], None,
                [8, 10, 6, 10, 8, 8, 10, 12, 9, 8, 9, 9, 7, 10, 6, 55])

    ws = wb.create_sheet("OUR_DETECTIONS")
    title(ws, "OUR DETECTIONS - every one, none merged (sorted on their metre axis)", "A1:K1")
    write_table(ws, 3, ["ID", "OUR_M", "ON_THEIR_AXIS", "POS_MM", "SIDE", "EDGE_MM",
                        "W_MM", "SCORE", "STATUS", "MATCHES_THEIR_SR", "IMAGE"], our_rows,
                [None, "0.000", "0.000", "0.0", None, "0.0", "0.0", SC_FMT, None, None, None], 9,
                [7, 10, 13, 9, 6, 9, 7, 9, 22, 16, 60], light=True)
    wb.save(out)

    # ---------------------------------------------------------------- console
    p = lambda v: "-" if v is None else "%.1f%%" % (v * 100)
    print("their spots judged %d (+%d outside view) | FOUND %d | LENGTH ONLY %d | MISSED %d"
          % (nT, len(outside), found, length_only, nT - found - length_only))
    print("our detections %d | compared %d | confirmed %d | no OS/DS %d | scrap/out %d"
          % (len(ours), n_cmp, n_conf, n_nolat, n_out))
    print("SPOT DETECTION ACCURACY %s | confirmation %s | F1 %s" % (p(rec), p(prec), p(f1)))
    print("alignment: %s" % align)
    if all_sc:
        f = (lambda v: "-" if v is None else ("%.1f%%" % (v * 100) if SC_FMT == "0.0%" else "%.2f" % v))
        print("score: confirmed %s | no counterpart %s | best matches %s (lowest %s)"
              % (f(mean(sc_conf)), f(mean(sc_none)), f(mean(best_sc)),
                 f(min(best_sc) if best_sc else None)))
    if zero_fit is not None:
        print("fitted: their 0 m = our %.3f m, K = %.4f" % (zero_fit, k_fit))
    for w in warns:
        print("WARNING: " + w)
    print("written: %s" % out)


if __name__ == "__main__":
    main()